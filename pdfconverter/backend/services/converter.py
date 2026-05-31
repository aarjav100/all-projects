import json
import csv
import io
import subprocess
import tempfile
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
import nbformat
from pptx import Presentation
from pptx.util import Inches, Pt

from storage import save_output


def convert_pdf(
    file_path: str,
    job_id: str,
    target_format: str,
    extracted_text: str = "",
    tables_json: str = "[]",
) -> str:
    """Convert a PDF to the target format. Returns output file path."""
    tables = json.loads(tables_json) if tables_json else []

    converters = {
        "txt": _to_txt,
        "md": _to_md,
        "html": _to_html,
        "docx": _to_docx,
        "ipynb": _to_ipynb,
        "csv": _to_csv,
        "pptx": _to_pptx,
    }

    fn = converters.get(target_format)
    if fn is None:
        raise ValueError(f"Unsupported format: {target_format}")

    return fn(file_path, job_id, extracted_text, tables)


def _to_txt(file_path, job_id, text, tables) -> str:
    data = text.encode("utf-8")
    return save_output(data, job_id, ".txt")


def _to_md(file_path, job_id, text, tables) -> str:
    lines = ["# Document\n"]
    for para in text.split("\n\n"):
        if para.strip():
            lines.append(para.strip())
            lines.append("")

    if tables:
        lines.append("\n## Tables\n")
        for tbl in tables:
            lines.append(f"### {tbl.get('title', 'Table')}\n")
            headers = tbl.get("headers", [])
            rows = tbl.get("rows", [])
            if headers:
                lines.append("| " + " | ".join(headers) + " |")
                lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
            for row in rows:
                lines.append("| " + " | ".join(row) + " |")
            lines.append("")

    data = "\n".join(lines).encode("utf-8")
    return save_output(data, job_id, ".md")


def _to_html(file_path, job_id, text, tables) -> str:
    paragraphs = "".join(
        f"<p>{p.strip()}</p>\n"
        for p in text.split("\n\n")
        if p.strip()
    )

    tables_html = ""
    for tbl in tables:
        headers = tbl.get("headers", [])
        rows = tbl.get("rows", [])
        header_html = "".join(f"<th>{h}</th>" for h in headers)
        rows_html = "".join(
            "<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>"
            for row in rows
        )
        tables_html += f"""
        <div class="table-wrap">
          <h3>{tbl.get('title', 'Table')}</h3>
          <table><thead><tr>{header_html}</tr></thead><tbody>{rows_html}</tbody></table>
        </div>"""

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>Converted Document</title>
  <style>
    body {{ font-family: system-ui, sans-serif; max-width: 800px; margin: 2rem auto; color: #1f2937; }}
    p {{ line-height: 1.7; margin-bottom: 1rem; }}
    table {{ border-collapse: collapse; width: 100%; margin: 1rem 0; }}
    th, td {{ border: 1px solid #e5e7eb; padding: 0.5rem 0.75rem; text-align: left; }}
    th {{ background: #f3f4f6; }}
  </style>
</head>
<body>
  <h1>Converted Document</h1>
  <section>{paragraphs}</section>
  {tables_html}
</body>
</html>"""

    return save_output(html.encode("utf-8"), job_id, ".html")


def _to_docx(file_path, job_id, text, tables) -> str:
    doc = Document()
    doc.add_heading("Converted Document", 0)

    for para in text.split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    if tables:
        doc.add_heading("Tables", level=1)
        for tbl_data in tables:
            headers = tbl_data.get("headers", [])
            rows = tbl_data.get("rows", [])
            if not headers and not rows:
                continue
            doc.add_heading(tbl_data.get("title", "Table"), level=2)
            all_rows = ([headers] if headers else []) + rows
            if all_rows:
                table = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
                table.style = "Table Grid"
                for i, row_data in enumerate(all_rows):
                    for j, cell_val in enumerate(row_data):
                        table.rows[i].cells[j].text = str(cell_val)

    buf = io.BytesIO()
    doc.save(buf)
    return save_output(buf.getvalue(), job_id, ".docx")


def _to_ipynb(file_path, job_id, text, tables) -> str:
    nb = nbformat.v4.new_notebook()
    cells = []

    cells.append(nbformat.v4.new_markdown_cell("# Converted PDF Document\n"))

    for para in text.split("\n\n"):
        if para.strip():
            cells.append(nbformat.v4.new_markdown_cell(para.strip()))

    if tables:
        cells.append(nbformat.v4.new_markdown_cell("## Tables\n"))
        for tbl_data in tables:
            headers = tbl_data.get("headers", [])
            rows = tbl_data.get("rows", [])
            if headers or rows:
                code = f"import pandas as pd\n"
                code += f"df = pd.DataFrame({json.dumps(rows)}, columns={json.dumps(headers)})\n"
                code += "df"
                cells.append(nbformat.v4.new_code_cell(code))

    nb.cells = cells
    data = nbformat.writes(nb).encode("utf-8")
    return save_output(data, job_id, ".ipynb")


def _to_csv(file_path, job_id, text, tables) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)

    if tables:
        for tbl_data in tables:
            writer.writerow([tbl_data.get("title", "Table")])
            headers = tbl_data.get("headers", [])
            rows = tbl_data.get("rows", [])
            if headers:
                writer.writerow(headers)
            writer.writerows(rows)
            writer.writerow([])
    else:
        # Fall back to plain text rows
        writer.writerow(["Content"])
        for line in text.split("\n"):
            if line.strip():
                writer.writerow([line.strip()])

    return save_output(buf.getvalue().encode("utf-8"), job_id, ".csv")


def _to_pptx(file_path, job_id, text, tables) -> str:
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Converted Document"
    slide.placeholders[1].text = "Generated by PDF Studio"

    # Content slides (one per paragraph block, max 10)
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()][:10]
    for para in paragraphs:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = para[:80]
        tf = slide.placeholders[1].text_frame
        tf.text = para[:500]

    # Table slides
    for tbl_data in tables[:5]:
        slide = prs.slides.add_slide(blank_layout)
        left = Inches(0.5)
        top = Inches(0.5)
        headers = tbl_data.get("headers", [])
        rows = tbl_data.get("rows", [])[:10]
        all_rows = ([headers] if headers else []) + rows
        if all_rows and all_rows[0]:
            tbl = slide.shapes.add_table(
                len(all_rows), len(all_rows[0]), left, top, Inches(9), Inches(0.4 * len(all_rows))
            ).table
            for i, row_data in enumerate(all_rows):
                for j, cell_val in enumerate(row_data):
                    tbl.cell(i, j).text = str(cell_val)

    buf = io.BytesIO()
    prs.save(buf)
    return save_output(buf.getvalue(), job_id, ".pptx")
